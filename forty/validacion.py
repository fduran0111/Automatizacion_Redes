from openpyxl import load_workbook
from servicios import cargar_catalogo
from seguridad import cargar_puerto_seguridad, evaluar
import ipaddress


def celdas_vacias(valor):
    if valor is None:
        return True
    elif str(valor).strip() == "":
        return True
    return False


def leer_solicitudes(ruta_excel):
    
    catalogo = cargar_catalogo()
    puerto_seguridad = cargar_puerto_seguridad()
    libro = load_workbook(ruta_excel,data_only=True)
    hoja = libro["CONFIGURADO"]
    solicitudes = []

    # Identificamos que fila es cada una(origen,destino,puerto)
    for fila in range(4, hoja.max_row + 1):
        origen = hoja.cell(fila,5).value # columna E
        destino = hoja.cell(fila,6).value # columna F
        puertos = hoja.cell(fila,7).value # columna G

        if celdas_vacias(origen) and celdas_vacias(destino) and celdas_vacias(puertos):
            continue
        solicitud = validar_solicitud(fila, origen, destino, puertos, catalogo, puerto_seguridad)    #valida la fila
        solicitudes.append(solicitud)  #se guarda al final de la fila
    return solicitudes #devuelve la lista completa

def validar_direcciones(valor, tipo):

    # Validamos si la celda esta vacia y le enviamos mensaje de lo que falta
    if celdas_vacias(valor):
        return [], [f"{tipo}: falta la direccion de {tipo.lower()}"]
    original = str(valor).strip()
    texto = original.replace(","," ").replace(";"," ")
    candidatos = texto.split()

    direcciones = [] #Candidatos que SI son ip o red
    sobrantes = [] #Candidatos que NO lo son

    #revisamos los candidatos si tienen mascara o no
    for candidato in candidatos:
        try:
            if "/" in candidato:
                direccion = ipaddress.ip_network(candidato, strict=False)
            else:
                direccion = ipaddress.ip_address(candidato)
            direcciones.append(str(direccion))
        except ValueError:
            sobrantes.append(candidato)

    #Si ninguna IP es valida
    if not direcciones:
        return [], [f"{tipo}: no es una IP/red reconocida, revisar si es un objeto: {original}"]

    # 5. Hay IPs, pero sobró texto
    errores = []
    if sobrantes:
        errores.append(f"{tipo}: texto adicional que requiere revisión: " + " ".join(sobrantes))

    return direcciones,errores

def validar_solicitud(fila, origen, destino, puertos, catalogo, puerto_seguridad):

    origenes, errores_origen = validar_direcciones(origen, "Origen")
    destinos, errores_destino = validar_direcciones(destino, "Destino")
    servicios, sin_servicio, errores_servicios = validar_servicios(puertos, catalogo)
    errores_seguridad, advertencias = evaluar(origenes, destinos, servicios, catalogo.datos, puerto_seguridad)
    errores = errores_origen + errores_destino + errores_servicios + errores_seguridad

    if errores:
        estado = "REVISAR"
    else:
        estado = "VALIDADA"

    return {
        "fila_excel": fila,
        "origenes": origenes,
        "destinos": destinos,
        "servicios": servicios,
        "sin_servicio": sin_servicio,
        "estado": estado,
        "errores": errores,
        "advertencias": advertencias,
    }

def validar_servicios(valor, catalogo):

    # Validamos si la celda está vacía
    if celdas_vacias(valor):
        return [], [], ["Falta el puerto o servicio"]

    texto = str(valor).strip()
    servicios = []       # nombres de los servicios que existen en el Forti

    # Primero los nombres con espacios ("Web Access"), que el split() cortaría
    for nombre in sorted(catalogo.con_espacios, key=len, reverse=True):
        posicion = texto.upper().find(nombre.upper())    # -1 si no lo encuentra
        while posicion != -1:
            if nombre not in servicios:
                servicios.append(nombre)
            # se borra del texto para que el split() no lo vea
            texto = texto[:posicion] + " " + texto[posicion + len(nombre):]
            posicion = texto.upper().find(nombre.upper())  # ¿aparece otra vez?

    # Separar en elementos (espacios, saltos de línea, comas y punto y coma)
    texto = texto.replace(",", " ").replace(";", " ")
    elementos = texto.split()

    sin_servicio = []    # puertos sin servicio en el Forti
    errores = []

    # Revisar cada elemento: n° de puerto o nombre de servicio
    for elemento in elementos:

        # Separar número y protocolo si viene como "514/udp" o "udp/514"
        numero = elemento
        protocolo = None

        if "/" in elemento:
            partes = elemento.lower().split("/")
            if len(partes) == 2 and partes[0] in ("tcp", "udp"):
                protocolo, numero = partes          # "udp/514"
            elif len(partes) == 2 and partes[1] in ("tcp", "udp"):
                numero, protocolo = partes          # "514/udp"

        # Número de puerto (con o sin protocolo)
        if numero.isdigit():
            puerto = int(numero)

            if puerto < 1 or puerto > 65535:
                errores.append(f"Puerto fuera del rango 1-65535: {elemento}")
                continue

            # Preguntar al catálogo
            nombre, error = catalogo.por_puerto(puerto, protocolo)

            if error:                                  # ambiguo
                errores.append(error)
            elif nombre:                               # encontró el servicio
                if nombre not in servicios:
                    servicios.append(nombre)
            else:                                      # no existe en el Forti
                if protocolo:
                    pedido = f"{protocolo}/{puerto}"
                else:
                    pedido = str(puerto)
                if pedido not in sin_servicio:
                    sin_servicio.append(pedido)
                errores.append(f"Falta el servicio en el Forti para el puerto {pedido}")

            continue

        # Nombre de servicio (ej. "RDP"); por_nombre ignora mayúsculas
        nombre = catalogo.por_nombre(elemento)

        if nombre:
            if nombre not in servicios:
                servicios.append(nombre)
        else:
            errores.append(f"El servicio no existe en el Forti: {elemento}")

    return servicios, sin_servicio, errores


def mostrar_resultados(solicitudes):

    validadas = 0
    con_warnings = 0

    # Detalle de cada fila
    for s in solicitudes:

        print(f"\nFila {s['fila_excel']} -> {s['estado']}")
        print("  Orígenes:", s["origenes"])
        print("  Destinos:", s["destinos"])
        print("  Servicios:", s["servicios"])

        for error in s["errores"]:
            print("  ERROR -", error)

        for advertencia in s["advertencias"]:
            print("  WARNING -", advertencia)

        # Contadores
        if s["estado"] == "VALIDADA":
            validadas += 1
        if s["advertencias"]:
            con_warnings += 1

    # Resumen final
    total = len(solicitudes)
    print("\nTotal de solicitudes:", total)
    print("Validadas:", validadas)
    print("Por revisar:", total - validadas)
    print("Con warnings:", con_warnings)

    # Resumen de puertos que faltan en el Forti
    mostrar_servicios_faltantes(solicitudes)

def mostrar_servicios_faltantes(solicitudes):

    # Juntar las filas de cada puerto: {"1433": ["4", "5", ...], "587": ["9"]}
    faltantes = {}

    for aux in solicitudes:
        for pedido in aux["sin_servicio"]:
            if pedido not in faltantes:
                faltantes[pedido] = []
            faltantes[pedido].append(str(aux["fila_excel"]))

    # Si no falta ningún servicio, no se muestra nada
    if not faltantes:
        return

    print("\n--- SERVICIOS QUE FALTAN EN EL FORTI ---")

    for pedido in faltantes:
        filas = ", ".join(faltantes[pedido])
        print(f"  Puerto {pedido} -> filas {filas}")

def direcciones_unicas(solicitudes):
    direcciones = set()

    for aux in solicitudes:

        # Solo las validadas: no tiene sentido crear objetos para filas en REVISAR
        if aux["estado"] != "VALIDADA":
            continue

        direcciones.update(aux["origenes"])
        direcciones.update(aux["destinos"])

    return sorted(direcciones)
