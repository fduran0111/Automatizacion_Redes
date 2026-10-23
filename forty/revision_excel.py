from validacion import leer_solicitudes
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.comments import Comment
from datetime import datetime
from pathlib import Path

# costantes de estilos para trabajar el excel
LETRA_NEGRITA = Font(bold=True)
LINEA = Side(style="thin")
BORDE = Border(left=LINEA, right=LINEA, top=LINEA, bottom=LINEA)
ROJO = PatternFill("solid", fgColor="F8CBAD")
VERDE = PatternFill("solid", fgColor="C6EFCE")   
AMARILLO = PatternFill("solid", fgColor="FFE699")
TITULOS = PatternFill("solid", fgColor="FFDE21")

def comentario_revision(celda,mensaje):
    #Validamos si tiene comentario, si tiene se "\n" abajo
    if celda.comment:
        texto = celda.comment.text + "\n" + mensaje 
    else:
        texto = mensaje

    comentario = Comment(texto, "Automatizacion")
    comentario.width = 300
    comentario.height = 100
    celda.comment = comentario


def columna_del_error(mensaje):
    #Identificamos el mensaje segun como empieza con esa frace para retornar cual es su columna
    if mensaje.startswith("Origen:"):
        return 5
    elif mensaje.startswith("Destino:") or "destinos" in mensaje:
        return 6
    else:
        return 7
    
def marcar_celdas_error(solicitudes,hoja):
    for aux in solicitudes:
        fila = aux["fila_excel"]


        for warning in aux["advertencias"]:
            columna = columna_del_error(warning)
            celda = hoja.cell(fila, columna)
            celda.fill = AMARILLO
            comentario_revision(celda, "WARNING: " + warning)

        for error in aux["errores"]:
            columna = columna_del_error(error)
            celda = hoja.cell(fila, columna)
            celda.fill = ROJO
            comentario_revision(celda, "ERROR: " + error)

def aplicar_bordes(hoja, fila_inicio, fila_fin):
    for fila in range(fila_inicio, fila_fin + 1):  
        for columna in (14, 15):                        
            hoja.cell(fila, columna).border = BORDE

def color_revision(solicitudes, hoja):
    #Aplicamos parametros en los titulos(colores,letras negras)
    titulo_estado = hoja.cell(3, 14, "ESTADOS")
    titulo_estado.fill = TITULOS
    titulo_estado.font = LETRA_NEGRITA

    titulo_revision = hoja.cell(3, 15, "REVISION")  
    titulo_revision.fill = TITULOS
    titulo_revision.font = LETRA_NEGRITA

    #Asignamos colores segun sus estados
    for aux in solicitudes:
        fila = aux["fila_excel"]
        if aux["estado"] != "VALIDADA":              
            texto = "REVISAR"
            color = ROJO
        elif aux["advertencias"]:                       
            texto = "WARNING"
            color = AMARILLO
        else:                                          
            texto = "OK"
            color = VERDE
        celda_color = hoja.cell(fila, 14, texto)
        celda_color.fill = color
        celda_color.font = LETRA_NEGRITA


def revision_errores(solicitudes, hoja):
    #Alargamos la columna "O" para ingresar texto
    hoja.column_dimensions["O"].width = 70

    #Comentamos segun su estado del error
    for aux in solicitudes:
        fila = aux["fila_excel"]
        mensajes = []
        for error in aux["errores"]:
            mensajes.append("ERROR: "+ error)

        for warning in aux["advertencias"]:
            mensajes.append("WARNING: "+ warning)

        texto = "\n".join(mensajes) 
        celda = hoja.cell(fila, 15, texto)
        celda.alignment = Alignment(wrap_text=True, horizontal="left")

             
def generar_excel(ruta_excel, solicitudes):
    
    libro = load_workbook(ruta_excel)
    hoja = libro["CONFIGURADO"]

    #SE CREA TODA LA COLUMNA DE ESTADOS CON SUS COLORES CORRESPONDIENTES
    color_revision(solicitudes, hoja)

    # SE CREA TODAS LAS COLUMNA DE REVISION  CON SUS ERRORES CORRESPONDIENTES
    revision_errores(solicitudes, hoja)

    # BORDES DE LA TABLAS
    ultima_fila = solicitudes[-1]["fila_excel"]
    aplicar_bordes(hoja, 3, ultima_fila)

    # Pintar la celda con el problema
    marcar_celdas_error(solicitudes, hoja)

    # salida del nuevo excel revisado
    ruta = Path(ruta_excel)
    fecha = datetime.now().strftime("%d_%m") 
    salida = f"{ruta.stem}_revisado_{fecha}{ruta.suffix}"
    libro.save(salida)

    return salida

